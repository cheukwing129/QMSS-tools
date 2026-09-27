#!/usr/bin/env python3
"""Convert the school's annual teacher timetable workbook to public web data.

Expected source layout: teacher name in A3; Day A–F in B:G; teaching periods
in rows 7, 8, 10, 11, 12, 14, 15, and 16. Requires openpyxl.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

DAYS = ["A", "B", "C", "D", "E", "F"]
PERIOD_ROWS = [7, 8, 10, 11, 12, 14, 15, 16]
CLASS_PATTERN = re.compile(r"(?<![A-Za-z0-9])([1-6][KPTW])(?=$|[\s,;/])")
CLASS_SORT = "KPTW"


def class_sort_key(class_code: str) -> tuple[int, int]:
    return int(class_code[0]), CLASS_SORT.index(class_code[1])


def display_name(full_name: str, code: str) -> str:
    if "," in full_name:
        chinese_name = re.sub(r"\s*\([^)]*\)\s*$", "", full_name.split(",", 1)[1]).strip()
    else:
        chinese_name = full_name
    if not re.search(r"[\u3400-\u9fff]", chinese_name):
        chinese_name = full_name
    return f"{chinese_name}（{code}）"


def convert(workbook_path: Path, academic_year: str) -> dict:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    teachers = []
    all_classes: set[str] = set()

    for sheet in workbook.worksheets:
        code = sheet.title.strip()
        full_name = str(sheet["A3"].value or "").strip()
        slot_groups: dict[tuple[str, int], dict[str, set[str] | list[str]]] = {}

        for column, day in enumerate(DAYS, start=2):
            for period, row in enumerate(PERIOD_ROWS, start=1):
                cell_value = sheet.cell(row, column).value
                if cell_value is None:
                    continue

                for line in str(cell_value).splitlines():
                    raw = " ".join(line.split()).strip()
                    if not raw:
                        continue
                    classes = list(dict.fromkeys(CLASS_PATTERN.findall(raw)))
                    if not classes:
                        continue

                    all_classes.update(classes)
                    subject = " ".join(CLASS_PATTERN.sub(" ", raw).split()).strip(" ,;/") or "（科目／地點未列明）"
                    group = slot_groups.setdefault((day, period), {"classes": set(), "subjects": []})
                    group["classes"].update(classes)
                    if subject not in group["subjects"]:
                        group["subjects"].append(subject)

        lessons = []
        for (day, period), group in sorted(slot_groups.items(), key=lambda item: (DAYS.index(item[0][0]), item[0][1])):
            lessons.append({
                "day": day,
                "period": period,
                "classes": sorted(group["classes"], key=class_sort_key),
                "subject": "／".join(group["subjects"]),
            })

        teachers.append({
            "code": code,
            "name": full_name,
            "displayName": display_name(full_name, code),
            "lessons": lessons,
        })

    return {
        "academicYear": academic_year,
        "sourceFile": workbook_path.name,
        "days": DAYS,
        "periods": list(range(1, 9)),
        "classes": sorted(all_classes, key=class_sort_key),
        "teachers": sorted(teachers, key=lambda item: (item["displayName"], item["code"])),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="年度教師課表 .xlsx 檔案")
    parser.add_argument("--year", required=True, help="學年，例如 2027–2028")
    parser.add_argument("--out", type=Path, default=Path("data/timetable.json"), help="輸出 JSON 路徑")
    args = parser.parse_args()

    timetable = convert(args.workbook, args.year)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(timetable, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lesson_count = sum(len(teacher["lessons"]) for teacher in timetable["teachers"])
    print(f"已產生 {args.out}: {len(timetable['teachers'])} 位老師、{len(timetable['classes'])} 個班別、{lesson_count} 節課堂")


if __name__ == "__main__":
    main()
